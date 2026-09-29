"""
parser.py — Multi-format document parser.

Supports: PDF · DOCX · XLSX · TXT · CSV · JSON · XML
"""

import json
from pathlib import Path
from typing import Any, Dict, Tuple, Union

from backend.core.utils import logger


class FormatParser:
    """
    Parse cybersecurity framework documents from multiple file formats.

    Returns (content, format_name) where content is either a plain
    string (PDF/DOCX/XLSX/TXT/CSV) or a dict (JSON/XML).
    """

    SUPPORTED = {".pdf", ".docx", ".xlsx", ".txt", ".csv", ".json", ".xml"}

    # ─────────────────────────────────────────
    # Public entry point
    # ─────────────────────────────────────────

    def parse(
        self, file_path: Union[str, Path]
    ) -> Tuple[Union[str, Dict], str]:
        """
        Parse *file_path* and return (content, format_string).

        Raises
        ------
        FileNotFoundError  – file does not exist
        ValueError         – extension not supported
        RuntimeError       – parser-level error with context
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED:
            raise ValueError(
                f"Unsupported format '{ext}'. "
                f"Supported: {', '.join(sorted(self.SUPPORTED))}"
            )

        _dispatch = {
            ".pdf":  self._parse_pdf,
            ".docx": self._parse_docx,
            ".xlsx": self._parse_xlsx,
            ".txt":  self._parse_txt,
            ".csv":  self._parse_csv,
            ".json": self._parse_json,
            ".xml":  self._parse_xml,
        }

        fmt = ext.lstrip(".")
        logger.info("Parsing %s file: %s", fmt.upper(), path.name)

        try:
            content = _dispatch[ext](path)
        except Exception as exc:
            raise RuntimeError(
                f"Failed to parse {path.name} ({fmt}): {exc}"
            ) from exc

        logger.info(
            "Parsed %s → %s chars",
            path.name,
            len(str(content)),
        )
        return content, fmt

    # ─────────────────────────────────────────
    # Format-specific parsers
    # ─────────────────────────────────────────

    def _parse_pdf(self, path: Path) -> str:
        """Extract text from all pages of a PDF."""
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)

    def _parse_docx(self, path: Path) -> str:
        """Extract paragraph text from a Word document."""
        from docx import Document

        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)

    def _parse_xlsx(self, path: Path) -> str:
        """
        Extract cell values from all sheets of an Excel workbook.
        Each row becomes a pipe-delimited string.
        """
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows.append(f"[Sheet: {sheet_name}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))

        return "\n".join(rows)

    def _parse_txt(self, path: Path) -> str:
        """Read plain-text file."""
        return path.read_text(encoding="utf-8", errors="replace")

    def _parse_csv(self, path: Path) -> str:
        """Load CSV via pandas and return string representation."""
        import pandas as pd

        df = pd.read_csv(path, dtype=str)
        return df.fillna("").to_string(index=False)

    def _parse_json(self, path: Path) -> Dict[str, Any]:
        """Load and return parsed JSON."""
        with open(path, "r", encoding="utf-8") as fh:
            return json.load(fh)

    def _parse_xml(self, path: Path) -> Dict[str, Any]:
        """Parse XML into an ordered dict via xmltodict."""
        import xmltodict

        with open(path, "r", encoding="utf-8") as fh:
            return xmltodict.parse(fh.read())


# ─────────────────────────────────────────
# Evidence-only parser (TXT / CSV / JSON)
# ─────────────────────────────────────────

class EvidenceParser:
    """
    Lightweight parser restricted to evidence file types:
    TXT, CSV, JSON.
    """

    SUPPORTED = {".txt", ".csv", ".json"}

    def parse(self, file_path: Union[str, Path]) -> str:
        """Return evidence content as a plain string."""
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"Evidence file not found: {path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED:
            raise ValueError(f"Evidence format '{ext}' not supported.")

        logger.info("Parsing evidence: %s", path.name)

        if ext == ".txt":
            return path.read_text(encoding="utf-8", errors="replace")

        if ext == ".csv":
            import pandas as pd
            df = pd.read_csv(path, dtype=str)
            return df.fillna("").to_string(index=False)

        if ext == ".json":
            with open(path, "r", encoding="utf-8") as fh:
                data = json.load(fh)
            # Flatten to string for LLM consumption
            return json.dumps(data, indent=2, ensure_ascii=False)

        return ""  # unreachable
