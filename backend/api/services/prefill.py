"""Self-contained evidence extraction + LLM pre-fill helpers for the API layer.

Historically the API imported these from ``backend.legacy_web.app`` — a heavy
Flask module that pulls in flask / werkzeug / pypdf / docx / chromadb, none of
which are in ``backend/api/requirements.txt``. In an API-only deployment that
import fails, so evidence text was never extracted (no ``.preview.txt`` sidecars
were written) and pre-fill silently produced zero answers.

This module depends only on packages the API already ships with: PyMuPDF for
PDF text and ``backend.core.utils`` (requests-based) for the LLM call.
"""
import base64
import json
import os
import re
import requests
from pathlib import Path

from backend.core.utils import LLMClient, parse_llm_json, logger


# ─────────────────────────────────────────────
# EVIDENCE TEXT EXTRACTION
# ─────────────────────────────────────────────

_TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".xml", ".log",
                  ".yaml", ".yml", ".html", ".htm", ".rst", ".ini", ".cfg"}

_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}

_MIME_MAP = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png"}

_IMAGE_VISION_MODEL = "meta-llama/llama-4-scout-17b-16e-instruct"

_IMAGE_EXTRACT_PROMPT = (
    "You are a security compliance analyst. This image is an evidence document or "
    "screenshot submitted for a cybersecurity compliance assessment.\n\n"
    "Carefully examine the image and extract ALL visible content:\n"
    "- Text, labels, settings, configurations, and values\n"
    "- System names, usernames, timestamps, and dates\n"
    "- Policy statements, controls, or procedures\n"
    "- Chart or table data with specific values\n"
    "- Security-relevant details: access controls, logs, alerts, configurations\n\n"
    "Write a detailed, factual description of everything visible. Be specific and "
    "thorough — your output is used as text evidence for compliance analysis. "
    "Do not infer or add information not visible in the image."
)


def _extract_image_text(path: Path) -> str:
    """Send an image to Groq llama-4-scout vision and return a text description."""
    api_key = os.getenv("GROQ_API_KEY", "")
    if not api_key:
        logger.warning("GROQ_API_KEY not set — cannot extract text from image %s", path.name)
        return ""
    try:
        mime = _MIME_MAP.get(path.suffix.lower(), "image/png")
        image_b64 = base64.b64encode(path.read_bytes()).decode("utf-8")
        payload = {
            "model": _IMAGE_VISION_MODEL,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": _IMAGE_EXTRACT_PROMPT},
                        {"type": "image_url", "image_url": {
                            "url": f"data:{mime};base64,{image_b64}"
                        }},
                    ],
                }
            ],
            "max_tokens": 1024,
            "temperature": 0.1,
        }
        resp = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            json=payload,
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            timeout=60,
        )
        resp.raise_for_status()
        return resp.json()["choices"][0]["message"]["content"]
    except Exception as exc:
        logger.warning("Image vision extraction failed for %s: %s", path.name, exc)
        return ""


def _extract_pdf(path: Path) -> str:
    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24
    except Exception:  # pragma: no cover - older PyMuPDF
        import fitz  # type: ignore
    parts: list[str] = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            parts.append(page.get_text())
    return "\n".join(parts)


def extract_evidence_text(file_path) -> str:
    """Extract a plain-text preview from a supported evidence file.

    Returns "" (never raises) for unsupported types or extraction errors so the
    caller can treat "no text" uniformly.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix in _IMAGE_SUFFIXES:
            return _extract_image_text(path)
        if suffix in _TEXT_SUFFIXES:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".docx":
            try:
                from docx import Document  # optional dependency
            except Exception:
                logger.warning("python-docx not installed; cannot read %s", path.name)
                return ""
            return "\n".join(p.text for p in Document(str(path)).paragraphs)
        if suffix in {".xlsx", ".xls"}:
            try:
                import openpyxl  # optional dependency
            except Exception:
                logger.warning("openpyxl not installed; cannot read %s", path.name)
                return ""
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            rows: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append("\t".join(cells))
                    if len(rows) >= 200:
                        break
                if len(rows) >= 200:
                    break
            wb.close()
            return "\n".join(rows)
    except Exception as exc:
        logger.warning("Evidence extraction failed for %s: %s", path, exc)
    return ""


# ─────────────────────────────────────────────
# LLM PRE-FILL
# ─────────────────────────────────────────────

PREFILL_SYSTEM_PROMPT = """\
You are a compliance analyst. You are given EVIDENCE documents collected from an
organisation and a list of assessment questions.

Answer EACH question STRICTLY and ONLY using facts stated in the provided
evidence. This is the absolute rule:
- Do NOT use outside knowledge, assumptions, or industry best practice.
- If the evidence does not contain enough information to answer a question, set
  "supported" to false and leave "answer" empty.
- Name the specific evidence file your answer is based on in "source".

Answer format by question_type:
- yes_no        : "Yes", "No", or "Partial"
- scale_1_5     : an integer 1-5 (1 = ad-hoc/initial, 5 = optimised)
- multi_choice  : EXACTLY one of the provided choices, copied verbatim
- maturity_rating: one of "Initial","Developing","Defined","Managed","Optimizing"
- free_text     : a concise answer (<= 80 words) stating only what the evidence shows

Be conservative: prefer "supported": false over guessing.

Return ONLY valid JSON (no markdown, no commentary):
{
  "answers": [
    {
      "question_id": "<id>",
      "answer": <type-appropriate value, or "" if unsupported>,
      "supported": true,
      "rationale": "<one sentence grounded in the evidence>",
      "source": "<evidence filename the answer is based on>"
    }
  ]
}
"""


def build_evidence_corpus(sections, max_chars: int = 60000) -> str:
    """Join (filename, text) evidence sections into a single capped corpus."""
    parts: list[str] = []
    used = 0
    for name, text in sections:
        if used >= max_chars:
            break
        header = f"\n### EVIDENCE FILE: {name}\n"
        budget = max_chars - used
        body = text if len(text) <= budget else (text[:budget] + "\n…[truncated]")
        chunk = header + body
        parts.append(chunk)
        used += len(chunk)
    return "".join(parts)


def select_evidence_for_batch(sections, batch_questions,
                              total_cap: int = 5000, per_file_cap: int = 2500) -> str:
    """Build a SMALL corpus containing only the evidence most relevant to this
    batch of questions, ranked by keyword overlap.

    Sending the full corpus to every batch blows past free-tier token-per-minute
    limits (HTTP 413 / 429). Each batch instead receives only its most relevant
    files, tightly capped, so every request stays small — e.g. the vendor batch
    pulls the vendor files, the cloud batch pulls the cloud files.
    """
    terms: dict = {}
    for q in batch_questions:
        blob = " ".join([
            str(q.get("question_text") or q.get("text") or ""),
            str(q.get("control_id") or ""),
            str(q.get("domain") or ""),
        ]).lower()
        for w in re.findall(r"[a-z][a-z0-9]{3,}", blob):
            terms[w] = terms.get(w, 0) + 1

    def _score(item):
        name, text = item
        hay = (name + " " + text[:4000]).lower()
        return sum(hay.count(w) * c for w, c in terms.items())

    ranked = sorted(sections, key=_score, reverse=True) if terms else list(sections)
    parts: list[str] = []
    used = 0
    for name, text in ranked:
        if used >= total_cap:
            break
        body = text[:per_file_cap]
        header = f"\n### EVIDENCE FILE: {name}\n"
        chunk = header + body
        parts.append(chunk)
        used += len(chunk)
    return "".join(parts)


def _prompt_type(question_type: str) -> str:
    """Map this app's UPPERCASE question types to the prompt's vocabulary."""
    t = (question_type or "").upper()
    if t == "YES_NO":
        return "yes_no"
    if t in ("SCALE_1_5", "SCALE"):
        return "scale_1_5"
    if t in ("MULTI_CHOICE", "MULTIPLE_CHOICE", "PRE_ASSESSMENT"):
        return "multi_choice"
    if t in ("MATURITY", "MATURITY_RATING"):
        return "maturity_rating"
    return "free_text"


def prefill_with_llm(questions, evidence_corpus, api_key, model, max_tokens=1024):
    """Ask the LLM to answer a batch of questions strictly from the evidence.

    Raises on a hard LLM failure (so the caller can surface it) rather than
    silently returning no answers.
    """
    payload_qs = [
        {
            "question_id": q.get("question_id"),
            "question_type": _prompt_type(q.get("question_type")),
            "question_text": q.get("question_text") or q.get("text", ""),
            "choices": q.get("choices") or [],
        }
        for q in questions
    ]
    user_prompt = (
        "EVIDENCE DOCUMENTS:\n"
        f"{evidence_corpus}\n\n"
        "QUESTIONS (answer each strictly from the evidence above):\n"
        f"{json.dumps(payload_qs, indent=2)}\n\n"
        "Return the JSON object now."
    )
    # Free-tier rate limits are PER MODEL; try the fast 8b model first, then the
    # configured model, then 70b — so one rate-limited model can't sink everything.
    model_chain: list = []
    for m in ["llama-3.1-8b-instant", model, "llama-3.3-70b-versatile"]:
        if m and m not in model_chain:
            model_chain.append(m)
    client = LLMClient(api_key, model_chain, temperature=0.0)
    raw = client.invoke({
        "system_prompt": PREFILL_SYSTEM_PROMPT,
        "user_prompt": user_prompt,
        "max_tokens": max_tokens,
    })
    parsed = parse_llm_json(raw)
    if isinstance(parsed, dict):
        return parsed.get("answers", []) or []
    if isinstance(parsed, list):
        return parsed
    return []


def map_prefill_answer(meta: dict, a: dict) -> dict | None:
    """Map a raw LLM answer to a form-ready {response, note, supported} dict.

    Crucially, ``response`` is emitted in the EXACT form the app's UI/scoring
    expect: UPPERCASE "YES"/"NO"/"PARTIAL", the verbatim choice for choice
    questions, and a bare digit for scales.
    """
    ptype = _prompt_type(meta.get("question_type"))
    supported = bool(a.get("supported"))
    raw = a.get("answer")
    raw_str = "" if raw is None else str(raw).strip()

    response = ""
    if supported and raw_str:
        if ptype == "yes_no":
            t = raw_str.lower()
            response = "YES" if t.startswith("y") else "PARTIAL" if t.startswith("p") else "NO" if t.startswith("n") else ""
        elif ptype == "scale_1_5":
            m = re.search(r"[1-5]", raw_str)
            response = m.group(0) if m else ""
        elif ptype == "multi_choice":
            choices = meta.get("choices") or []
            low = raw_str.lower()
            match = next((c for c in choices if c.strip().lower() == low), None)
            if not match:
                match = next((c for c in choices if low in c.strip().lower() or c.strip().lower() in low), None)
            if not match and choices:
                # The LLM often returns the maturity level as a number (e.g. "3"
                # or "level 3") because it reads the item's 0-3 rubric, instead of
                # the choice label. Map a valid 0-based index to its choice so the
                # stored value is the label the UI dropdown AND the scorer's anchor
                # resolver expect (a bare "3" matches neither and is scored as a
                # non-anchor, silently dropping the item to the fallback level).
                mnum = re.search(r"\d+", raw_str)
                if mnum:
                    idx = int(mnum.group(0))
                    if 0 <= idx < len(choices):
                        match = choices[idx]
            response = match or raw_str
        elif ptype == "maturity_rating":
            ladder = ["Initial", "Developing", "Defined", "Managed", "Optimizing"]
            response = next((m for m in ladder if m.lower() == raw_str.lower()), raw_str.capitalize())
        else:  # free_text and anything else
            response = raw_str

    rationale = str(a.get("rationale") or "").strip()
    source = str(a.get("source") or "").strip()
    if supported and response:
        note = "[AI pre-filled] " + (rationale or "Derived from submitted evidence.")
        if source:
            note += f" (Source: {source})"
    else:
        note = "[AI pre-fill] Insufficient evidence to answer this question."
    return {"response": response, "note": note, "supported": supported and bool(response)}
